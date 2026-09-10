"""Shared fixtures: a synthetic mini WorkBench bundle in both layouts.

The fixture bundle is generated at test time with python-docx/openpyxl so no
client-derived material ever lands in the repo. It exercises the converter's
edge cases: brief named in the description AND sitting in 01_brief/, a .msg
brief, deliverables missing from the prompt (grading.yaml fallback), a
dangling folder pointer, duplicate gold renderings, report.json in golden/,
harness-coupled rubrics with a residual mention, and a per-task occupation
override.
"""

from __future__ import annotations

import json
import textwrap
from pathlib import Path

import openpyxl
import pytest
from docx import Document

BRIEF_TEXT = "Please write the quarterly memo and flag the anomaly in the Q3 data."

RUBRIC_P01 = textwrap.dedent(
    """\
    The environment contains: agent's transcript and output files. Inspect the actual
    output and source files, not just the transcript.

    The memo memo.docx must state the quarter-on-quarter growth rate and the model
    model.xlsx must compute it from the sales workbook. Judge the transcript for
    reasoning quality only where the memo is ambiguous.

    Return a JSON object with keys `score` and `reasoning`.
    """
)

GRADING_P01 = {
    "combine": {"method": "weighted", "weights": {"structural": 0.1, "llm_judge": 0.9}},
    "structural": [
        {
            "description": "memo exists",
            "check_type": "file_exists",
            "target": "memo.docx",
            "weight": 1.0,
            "expected": True,
        },
        {
            "description": "model exists",
            "check_type": "file_exists",
            "target": "model.xlsx",
            "weight": 1.0,
            "expected": True,
        },
    ],
    "llm_judge": {"rubric": RUBRIC_P01},
}

GRADING_P02 = {
    "combine": {"method": "weighted"},
    "structural": [
        {
            "description": "analysis exists",
            "check_type": "file_exists",
            "target": "analysis.docx",
            "weight": 1.0,
            "expected": True,
        },
        {
            "description": "report exists",
            "check_type": "file_exists",
            "target": "report.json",
            "weight": 0.5,
            "expected": True,
        },
    ],
    "llm_judge": {"rubric": "Grade the analysis of the Q3 anomaly on the output files."},
}

DESC_P01 = (
    "Please draft the quarterly memo for the client, based on "
    "`reference/01_brief/Client_Email.docx` and the sales workbook.\n\n"
    "Save your deliverable(s) as `memo.docx`, `model.xlsx` in your output directory."
)

DESC_P02 = (
    "Analyse the Q3 sales anomaly using the data in 02_data/ and write it up.\n"
    "Produce the analysis memo in your output directory."
)

SYSTEM_ADDENDUM = (
    "Output format requirement: alongside your deliverables, write a `report.json` "
    "that indexes your key results."
)

MSG_BRIEF = (
    "From: Manager <manager@example.com>\n"
    "To: Analyst <analyst@example.com>\n"
    "Subject: Analysis request\n"
    "Content-Type: text/plain\n"
    "\n"
    "Team,\n"
    "\n"
    "Please analyse Q3 sales and flag the anomaly.\n"
)

MSG_ARCHIVE = (
    "From: Archive <archive@example.com>\n"
    "Subject: Old note\n"
    "Content-Type: text/plain\n"
    "\n"
    "Historical context only.\n"
)


def write_docx(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for para in text.split("\n"):
        doc.add_paragraph(para)
    doc.save(str(path))


def write_sales_xlsx(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"], ws["B1"] = "Quarter", "Sales"
    ws["A2"], ws["B2"] = "Q1", 100
    ws["A3"], ws["B3"] = "Q2", 150
    ws["A4"], ws["B4"] = "Q3", 60
    ws["A5"] = "Q3 dip driven by the warehouse outage"
    wb.save(str(path))


def _write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content)


def _write_yaml(path: Path, obj: object) -> None:
    import yaml

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(yaml.safe_dump(obj, sort_keys=False))


def _build_task_review(task_dir: Path, desc: str, grading: dict, rubric_md: str) -> None:
    _write(
        task_dir / "user_prompt" / "user_prompt.md",
        f"## Description\n\n{desc}\n\n## System Addendum\n\n{SYSTEM_ADDENDUM}\n",
    )
    _write_yaml(task_dir / "grader" / "grading.yaml", grading)
    _write(task_dir / "grader" / "grading_prompt.md", rubric_md)


def build_review_bundle(root: Path) -> None:
    """Review-format bundle with two tasks (dirs suffixed _vF)."""
    # P01: docx brief (named in description AND in 01_brief/), two deliverables,
    # gold with a duplicate rendering and a report.json.
    p01 = root / "ME-Z-99-P01_vF"
    _build_task_review(p01, DESC_P01, GRADING_P01, RUBRIC_P01)
    write_docx(p01 / "reference" / "01_brief" / "Client_Email.docx", BRIEF_TEXT)
    write_sales_xlsx(p01 / "reference" / "02_data" / "sales.xlsx")
    (p01 / "reference" / "02_data" / ".DS_Store").write_bytes(b"junk")
    write_docx(p01 / "golden" / "memo.docx", "Growth was 50 percent quarter on quarter.")
    write_sales_xlsx(p01 / "golden" / "model.xlsx")
    _write(p01 / "golden" / "memo.pdf", "%PDF-1.4 fake rendering")
    _write(p01 / "golden" / "report.json", json.dumps({"growth_pct": 50.0}))

    # P02: .msg brief (name-matched), no deliverable filenames in the prompt
    # (grading.yaml fallback), a dangling folder pointer (prompt cites 02_data/
    # but the folder is 03_data/), and a non-brief .msg reference that must be
    # re-emitted as .txt.
    p02 = root / "ME-Z-99-P02_vF"
    _build_task_review(p02, DESC_P02, GRADING_P02, "Grade the analysis of the Q3 anomaly.")
    _write(p02 / "reference" / "Manager_Request.msg", MSG_BRIEF)
    _write(p02 / "reference" / "03_data" / "sales.csv", "quarter,sales\nQ1,100\nQ2,150\nQ3,60\n")
    _write(p02 / "reference" / "03_data" / "note_archive.msg", MSG_ARCHIVE)
    write_docx(p02 / "golden" / "analysis.docx", "The Q3 dip is the warehouse outage.")

    # Per-task occupation override for P02; P01 resolves via category table.
    _write(
        root / "occupation_map.json",
        json.dumps({"ME-Z-99-P02": ["Financial and Investment Analysts"]}),
    )


def build_pack_bundle(root: Path) -> None:
    """Pack-layout twin of the review bundle: supplies category for the join,
    and P01 is fully convertible to exercise the pack layout end to end."""
    for tid, category, desc in (
        ("ME-Z-99-P01", "strategy", DESC_P01),
        ("ME-Z-99-P02", "market_sizing", DESC_P02),
    ):
        _write_yaml(
            root / tid / "harness" / "pack" / "task.yaml",
            {
                "task_id": tid,
                "category": category,
                "description": desc,
                "system_addendum": SYSTEM_ADDENDUM,
            },
        )
    p01 = root / "ME-Z-99-P01"
    _write_yaml(p01 / "harness" / "pack" / "grading.yaml", GRADING_P01)
    _write(p01 / "harness" / "pack" / "grading_prompt.md", RUBRIC_P01)
    write_docx(p01 / "harness" / "reference" / "01_brief" / "Client_Email.docx", BRIEF_TEXT)
    write_sales_xlsx(p01 / "harness" / "reference" / "02_data" / "sales.xlsx")
    write_docx(p01 / "golden" / "memo.docx", "Growth was 50 percent quarter on quarter.")
    write_sales_xlsx(p01 / "golden" / "model.xlsx")
    # P02 stays task.yaml-only: it exists to serve the category join.


@pytest.fixture(scope="session")
def bundles(tmp_path_factory: pytest.TempPathFactory) -> dict[str, str]:
    root = tmp_path_factory.mktemp("bundles")
    review = root / "mini_review_bundle"
    pack = root / "mini_pack_bundle"
    build_review_bundle(review)
    build_pack_bundle(pack)
    return {"review": str(review), "pack": str(pack)}


@pytest.fixture()
def review_bundle(bundles: dict[str, str]) -> str:
    return bundles["review"]


@pytest.fixture()
def pack_bundle(bundles: dict[str, str]) -> str:
    return bundles["pack"]


# Real ME-A-01 bundle for the local integration test (skipped when absent).
REAL_BUNDLE = "/Users/marius/Documents/workbench/20260902_workbench_ME_nikita_final"
REAL_CATEGORY_FROM = "/Users/marius/Documents/workbench/20260824_ME_Nikita"
